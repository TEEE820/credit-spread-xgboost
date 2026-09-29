"""论文 Word 排版生成：解析 paper/论文草稿.md，按《华南师范大学学报》格式生成 docx。

标记语法：
  # 标题        -> 论文主标题（三号黑体居中）
  ## 标题       -> 一级标题（小四黑体居中，上下空一行）
  ### 标题      -> 二级标题（五号仿宋）
  #### 标题     -> 三级标题（五号楷体）
  > 作者：...    -> 作者行（小四楷体居中）
  > 摘要：...    -> 摘要（小五仿宋，首行缩进）
  > 关键词：...  -> 关键词（小五仿宋）
  > 中图分类号... -> 小五仿宋
  > REFn ...    -> 参考文献条目（小五宋体，悬挂缩进）
  > 英文题目：... -> 英文题目（小五宋体加粗）
  > ABSTRACT: ... -> 英文摘要（小五宋体）
  > KEYWORDS: ... -> 英文关键词（小五宋体）
  [TBL]csv|表注  -> 三线表 + 表注（小五黑体居中）
  [IMG]png|图注  -> 图片 + 图注（小五黑体居中）
  其他行        -> 正文（五号宋体，首行缩进2字符，1.5倍行距）
"""
import os
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.normpath(os.path.join(os.path.dirname(__file__), ".."))
MD = os.path.join(ROOT, "paper", "论文草稿.md")
OUT = os.path.join(ROOT, "paper", "中国信用利差驱动因素的非线性机制与样本外预测.docx")

doc = Document()

# ---------- 页面与默认样式 ----------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.54)
sec.left_margin = sec.right_margin = Cm(2.6)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(10.5)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def set_font(run, cn_font, size, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn_font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def para(text, cn_font="宋体", size=10.5, align=None, indent=None, bold=False,
         space_before=0, space_after=0, line=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, cn_font, size, bold=bold)
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    if indent is not None:
        pf.first_line_indent = Pt(indent)
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if line is not None:
        pf.line_spacing = line
    return p


def set_cell_borders(cell, top=None, bottom=None):
    """三线表边框：top/bottom 传 ('single', sz) 或 None。"""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, spec in (("top", top), ("bottom", bottom)):
        if spec is None:
            continue
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), spec[0])
        el.set(qn("w:sz"), str(spec[1]))
        el.set(qn("w:color"), "000000")


def add_three_line_table(df, caption):
    n_rows, n_cols = df.shape
    table = doc.add_table(rows=n_rows + 1, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 表头
    for j, col in enumerate(df.columns):
        cell = table.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(str(col))
        set_font(r, "宋体", 9, bold=False)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 数据行
    for i in range(n_rows):
        for j in range(n_cols):
            cell = table.rows[i + 1].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(df.iloc[i, j]))
            set_font(r, "宋体", 9)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 三线边框
    for j in range(n_cols):
        set_cell_borders(table.rows[0].cells[j], top=("single", 12), bottom=("single", 6))
        set_cell_borders(table.rows[-1].cells[j], bottom=("single", 12))
        for i in range(1, n_rows):
            set_cell_borders(table.rows[i].cells[j])
    # 表注
    para(caption, cn_font="黑体", size=9, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_before=3, space_after=10)


def add_figure(png, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(png, width=Cm(14.0))
    para(caption, cn_font="黑体", size=9, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_before=3, space_after=10)


def load_table(csv_path):
    name = os.path.basename(csv_path)
    if name == "table1_summary_stats.csv":
        df = pd.read_csv(csv_path, index_col=0)
        df.index.name = "统计量"
        df = df.round(4).reset_index()
    elif name == "table2_model_performance.csv":
        df = pd.read_csv(csv_path)
        df = df.rename(columns={"ols": "OLS", "random_forest": "随机森林", "xgboost": "XGBoost"})
        df = df.round(4)
    else:  # table3_rolling_metrics.csv
        df = pd.read_csv(csv_path)
        df["方向准确率"] = (df["方向准确率"] * 100).round(1).astype(str) + "%"
        df["RMSE改进%"] = df["RMSE改进%"].round(1).astype(str) + "%"
        for c in ("RMSE", "MAE", "随机游走RMSE"):
            df[c] = df[c].round(4)
    return df


def shade_paragraph(p, fill="F2F2F2"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def add_code_block(code_lines):
    """代码块：等宽字体小五号、浅灰底纹、单倍行距。"""
    for i, code in enumerate(code_lines):
        p = doc.add_paragraph()
        r = p.add_run(code if code else " ")
        set_font(r, "宋体", 9)
        r.font.name = "Consolas"
        pf = p.paragraph_format
        pf.line_spacing = 1.0
        pf.space_before = Pt(0)
        pf.space_after = Pt(0 if i < len(code_lines) - 1 else 8)
        pf.left_indent = Pt(18)
        shade_paragraph(p, "F2F2F2")


# ---------- 解析并生成 ----------
lines = open(MD, encoding="utf-8").read().splitlines()
ref_no = 0
in_code = False
code_buf = []
for ln in lines:
    s = ln.strip()
    if s.startswith("```"):
        if in_code:
            add_code_block(code_buf)
            code_buf = []
            in_code = False
        else:
            in_code = True
        continue
    if in_code:
        code_buf.append(ln)
        continue
    if not s:
        continue
    if s.startswith("[TBL]"):
        body = s[5:]
        csv_name, caption = body.split("|", 1)
        add_three_line_table(load_table(os.path.join(ROOT, csv_name)), caption.strip())
    elif s.startswith("[IMG]"):
        body = s[5:]
        png, caption = body.split("|", 1)
        add_figure(os.path.join(ROOT, png), caption.strip())
    elif s.startswith("#### "):
        para(s[5:], cn_font="楷体", size=10.5, bold=True, space_before=6, space_after=3)
    elif s.startswith("### "):
        para(s[4:], cn_font="仿宋", size=10.5, bold=True, space_before=6, space_after=3)
    elif s.startswith("## "):
        para(s[3:], cn_font="黑体", size=12, align=WD_ALIGN_PARAGRAPH.CENTER,
             space_before=10, space_after=10)
    elif s.startswith("# "):
        para(s[2:], cn_font="黑体", size=16, align=WD_ALIGN_PARAGRAPH.CENTER,
             space_before=6, space_after=14)
    elif s.startswith("> "):
        body = s[2:]
        if body.startswith("作者："):
            para(body, cn_font="楷体", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
        elif body.startswith("摘要："):
            p = doc.add_paragraph()
            r = p.add_run("摘要：")
            set_font(r, "仿宋", 9, bold=True)
            r2 = p.add_run(body[3:])
            set_font(r2, "仿宋", 9)
            p.paragraph_format.first_line_indent = Pt(18)
            p.paragraph_format.space_after = Pt(3)
        elif body.startswith("关键词："):
            p = doc.add_paragraph()
            r = p.add_run("关键词：")
            set_font(r, "仿宋", 9, bold=True)
            r2 = p.add_run(body[4:])
            set_font(r2, "仿宋", 9)
            p.paragraph_format.first_line_indent = Pt(18)
            p.paragraph_format.space_after = Pt(3)
        elif body.startswith("中图分类号"):
            para(body, cn_font="仿宋", size=9, space_after=10)
        elif body.startswith("REF"):
            ref_no += 1
            text = body[body.index(" ") + 1:]
            p = para(f"[{ref_no}] {text}", cn_font="宋体", size=9, space_after=2, line=1.25)
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-18)
        elif body.startswith("英文题目："):
            para(body[5:], cn_font="宋体", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                 space_before=14, space_after=6)
        elif body.startswith("ABSTRACT:"):
            p = doc.add_paragraph()
            r = p.add_run("ABSTRACT: ")
            set_font(r, "宋体", 9, bold=True)
            r2 = p.add_run(body[9:].strip())
            set_font(r2, "宋体", 9)
            p.paragraph_format.space_after = Pt(3)
        elif body.startswith("KEYWORDS:"):
            p = doc.add_paragraph()
            r = p.add_run("KEYWORDS: ")
            set_font(r, "宋体", 9, bold=True)
            r2 = p.add_run(body[9:].strip())
            set_font(r2, "宋体", 9)
    else:
        para(s, cn_font="宋体", size=10.5, indent=21, space_after=2, line=1.5)

doc.save(OUT)
print("saved:", OUT)
